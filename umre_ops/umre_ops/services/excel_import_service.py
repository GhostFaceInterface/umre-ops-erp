# Copyright (c) 2026, Sermed Turizm and contributors
# For license information, please see license.txt

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import uuid4

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate, now
from frappe.utils.background_jobs import is_job_enqueued
from openpyxl import load_workbook

from umre_ops.umre_ops.services import cost_engine

DOCTYPE_IMPORT = "Umre Excel Import"
DOCTYPE_UMRECI = "Umreci"
DOCTYPE_BOOKING = "Umre Booking"
DOCTYPE_REFERRAL = "Referral Source"
RETRYABLE_IMPORT_ERRORS = (frappe.db.InternalError, frappe.RetryBackgroundJobError)
READY_ROW_STATUSES = {"Create Ready", "Update Ready", "No-op"}

IMPORT_FIELDS = (
	"TC KİMLİK",
	"AD",
	"SOYAD",
	"CİNSİYET",
	"UYRUK",
	"DOĞUM TARİHİ",
	"GELDİĞİ İL",
	"ODA SAYISI",
	"TELEFON NUMARASI",
	"KİMDEN",
	"FİYAT",
	"ÖDEDİĞİ MİKTAR",
	"YOLCU STATÜSÜ",
)
STATUS_ALLOWLIST = {
	"HOCA",
	"UMRECI",
	"HOCA_ESI",
	"HOCA_COCUGU",
	"SIRKET_MUDURU",
	"SIRKET_MUDURU_ESI",
	"SIRKET_MUDURU_COCUGU",
	"FREE",
}
NATIONALITY_CODES = {
	"tc": "TC",
	"tr": "TC",
	"tur": "TC",
	"turkiye": "TC",
	"turk": "TC",
	"turkish": "TC",
}

TURKISH_PROVINCES = (
	"Adana",
	"Adıyaman",
	"Afyonkarahisar",
	"Ağrı",
	"Aksaray",
	"Amasya",
	"Ankara",
	"Antalya",
	"Ardahan",
	"Artvin",
	"Aydın",
	"Balıkesir",
	"Bartın",
	"Batman",
	"Bayburt",
	"Bilecik",
	"Bingöl",
	"Bitlis",
	"Bolu",
	"Burdur",
	"Bursa",
	"Çanakkale",
	"Çankırı",
	"Çorum",
	"Denizli",
	"Diyarbakır",
	"Düzce",
	"Edirne",
	"Elazığ",
	"Erzincan",
	"Erzurum",
	"Eskişehir",
	"Gaziantep",
	"Giresun",
	"Gümüşhane",
	"Hakkari",
	"Hatay",
	"Iğdır",
	"Isparta",
	"İstanbul",
	"İzmir",
	"Kahramanmaraş",
	"Karabük",
	"Karaman",
	"Kars",
	"Kastamonu",
	"Kayseri",
	"Kırıkkale",
	"Kırklareli",
	"Kırşehir",
	"Kilis",
	"Kocaeli",
	"Konya",
	"Kütahya",
	"Malatya",
	"Manisa",
	"Mardin",
	"Mersin",
	"Muğla",
	"Muş",
	"Nevşehir",
	"Niğde",
	"Ordu",
	"Osmaniye",
	"Rize",
	"Sakarya",
	"Samsun",
	"Siirt",
	"Sinop",
	"Sivas",
	"Şanlıurfa",
	"Şırnak",
	"Tekirdağ",
	"Tokat",
	"Trabzon",
	"Tunceli",
	"Uşak",
	"Van",
	"Yalova",
	"Yozgat",
	"Zonguldak",
)


@dataclass
class ImportSummary:
	total_rows: int = 0
	created_umreci: int = 0
	updated_umreci: int = 0
	created_bookings: int = 0
	updated_bookings: int = 0
	row_errors: int = 0


@dataclass(frozen=True)
class ImportSourceSnapshot:
	status: str
	import_file: str | None
	target_tour: str | None
	header_row: int
	mapping_signature: str


class ImportRowError(Exception):
	def __init__(self, message: str, staged_row: dict[str, Any]):
		super().__init__(message)
		self.staged_row = staged_row


def header_key(value: Any) -> str:
	text = re.sub(r"\s+", " ", "" if value is None else str(value)).strip()
	text = text.replace("ı", "i").replace("İ", "I")
	text = unicodedata.normalize("NFKD", text)
	return "".join(ch for ch in text if unicodedata.category(ch) != "Mn").casefold()


def preserve_excel_text(value: Any) -> str:
	if is_blank(value):
		return ""
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value).strip()


def is_blank(value: Any) -> bool:
	return value is None or (isinstance(value, str) and not value.strip())


def normalize_tc(value: Any) -> str | None:
	text = re.sub(r"\s+", "", preserve_excel_text(value)).upper()
	if not text:
		return None
	if text.endswith(".0") and text[:-2].isdigit():
		text = text[:-2]
	if text.isdigit() and len(text) == 10:
		text = text.zfill(11)
	return text


def sanitize_phone(value: Any) -> str | None:
	text = preserve_excel_text(value)
	return text or None


def parse_date(value: Any) -> str | None:
	if is_blank(value):
		return None
	if isinstance(value, datetime):
		return value.date().isoformat()
	if isinstance(value, date):
		return value.isoformat()
	try:
		return getdate(value).isoformat()
	except Exception:
		return None


def safe_float(value: Any) -> float:
	if is_blank(value):
		return 0.0
	if isinstance(value, bool):
		frappe.throw(_("Sayısal değer geçersiz: {0}").format(repr(value)))
	text = preserve_excel_text(value)
	if re.fullmatch(r"-?\d{1,3}(\.\d{3})*,\d+", text):
		text = text.replace(".", "").replace(",", ".")
	elif re.fullmatch(r"-?\d{1,3}(,\d{3})*\.\d+", text):
		text = text.replace(",", "")
	elif re.fullmatch(r"-?\d+(?:,\d+)?", text):
		text = text.replace(",", ".")
	elif not re.fullmatch(r"-?\d+(?:\.\d+)?", text):
		frappe.throw(_("Sayısal değer geçersiz: {0}").format(repr(value)))
	try:
		return float(Decimal(text))
	except (InvalidOperation, ValueError):
		frappe.throw(_("Sayısal değer geçersiz: {0}").format(repr(value)))


def map_oda_tipi(value: Any) -> str:
	match = re.fullmatch(
		r"\s*([1-4])(?:\s*(?:kişi(?:lik)?|kisi(?:lik)?))?\s*", preserve_excel_text(value), re.I
	)
	if not match:
		frappe.throw(_("ODA SAYISI yalnızca 1, 2, 3 veya 4 olabilir."))
	return f"{match.group(1)} Kişilik"


def normalize_cinsiyet(value: Any) -> str:
	key = header_key(value)
	if key == "mrs":
		return "MRS"
	if key == "mr":
		return "MR"
	frappe.throw(_("CİNSİYET MR veya MRS olmalıdır."))


def normalize_nationality(value: Any) -> str:
	text = preserve_excel_text(value)
	if not text:
		frappe.throw(_("UYRUK zorunludur."))
	key = header_key(text)
	code = NATIONALITY_CODES.get(key, text.upper())
	if not re.fullmatch(r"[A-Z]{1,3}", code):
		frappe.throw(_("UYRUK 1-3 harfli bir kod olmalıdır (ör. TC, D, GB)."))
	return code


def normalize_status(value: Any) -> str:
	status = re.sub(r"[\s-]+", "_", header_key(value)).upper()
	status = {
		"SIRKET_MUDURUNUN_ESI": "SIRKET_MUDURU_ESI",
		"SIRKET_MUDURUNUN_COCUGU": "SIRKET_MUDURU_COCUGU",
	}.get(status, status)
	if status not in STATUS_ALLOWLIST:
		frappe.throw(_("YOLCU STATÜSÜ geçersiz: {0}").format(preserve_excel_text(value)))
	return status


def _city_compare_key(value: Any) -> str:
	return header_key(str(value).replace("'", "").replace('"', ""))


CITY_LOOKUP = {_city_compare_key(city): city for city in TURKISH_PROVINCES}


def normalize_city(value: Any) -> str:
	city = CITY_LOOKUP.get(_city_compare_key(value))
	if not city:
		frappe.throw(_("İl adı eşleşmedi: {0}").format(repr(value)))
	return city


def split_cities(value: Any) -> tuple[str, str]:
	text = preserve_excel_text(value)
	if not text:
		frappe.throw(_("GELDİĞİ İL zorunludur."))
	parts = [part.strip() for part in re.split(r"\s*(?:-|/|→)\s*", text, maxsplit=1)]
	arrival = normalize_city(parts[0])
	return_city = normalize_city(parts[1]) if len(parts) == 2 and parts[1] else arrival
	return arrival, return_city


def _mapping_rows(import_doc) -> list[dict[str, str]]:
	return sorted(
		(
			{"target_field": row.target_field, "source_column": row.source_column}
			for row in import_doc.get("column_mappings") or []
		),
		key=lambda row: row["target_field"],
	)


def _mapping_signature(import_doc) -> str:
	return hashlib.sha256(
		json.dumps(_mapping_rows(import_doc), ensure_ascii=False, sort_keys=True).encode()
	).hexdigest()


def make_source_snapshot(import_doc) -> ImportSourceSnapshot:
	return ImportSourceSnapshot(
		import_doc.status,
		import_doc.import_file,
		import_doc.target_tour,
		cint(import_doc.header_row or 1),
		_mapping_signature(import_doc),
	)


def validate_dry_run_snapshot(snapshot: ImportSourceSnapshot, current_doc) -> None:
	if current_doc.status in {"Queued", "Processing", "Partially Completed", "Completed"}:
		frappe.throw(_("Bu aktarımın durumu değişti; kuru çalıştırma sonucu kaydedilmedi."))
	if make_source_snapshot(current_doc) != snapshot:
		frappe.throw(
			_("Dosya, hedef tur, başlık satırı, kolon eşlemesi veya durum değişti; sonuç kaydedilmedi.")
		)


def attempt_can_run(current_doc, attempt_id: str) -> bool:
	return current_doc.job_id == attempt_id and current_doc.status in {"Queued", "Processing"}


def _get_file_content(import_doc) -> bytes:
	return _get_file_content_by_url(import_doc.import_file)


def _get_file_content_by_url(file_url: str) -> bytes:
	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		frappe.throw(_("Yüklenen Excel dosyası bulunamadı."))
	file_doc = frappe.get_doc("File", file_name)
	file_doc.check_permission("read")
	content = file_doc.get_content()
	return content.encode() if isinstance(content, str) else bytes(content)


def _open_workbook(content: bytes):
	try:
		return load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
	except Exception:
		frappe.throw(_("Excel dosyası okunamadı. Geçerli bir .xlsx dosyası yükleyin."))


def _read_single_worksheet(content: bytes) -> list[list[Any]]:
	workbook = _open_workbook(content)
	try:
		if len(workbook.sheetnames) != 1:
			frappe.throw(_("İçe aktarım dosyası tam olarak bir fiziksel çalışma sayfası içermelidir."))
		return [list(row) for row in workbook[workbook.sheetnames[0]].iter_rows(values_only=True)]
	finally:
		workbook.close()


def inspect_headers(import_doc, content: bytes | None = None) -> list[str]:
	rows = _read_single_worksheet(content or _get_file_content(import_doc))
	header_row = cint(import_doc.header_row or 1)
	if header_row < 1 or header_row > len(rows):
		frappe.throw(_("Başlık satırı Excel dosyasının dışında."))
	return [preserve_excel_text(value) for value in rows[header_row - 1] if not is_blank(value)]


def inspect_headers_from_file(file_url: str, header_row: int = 1) -> list[str]:
	"""Inspect an uploaded workbook before its parent import document is saved."""
	probe = frappe._dict(import_file=file_url, header_row=cint(header_row or 1))
	return inspect_headers(probe, _get_file_content_by_url(file_url))


def suggest_column_mappings(headers: list[str]) -> list[dict[str, str]]:
	"""Suggest normalized exact matches; unmatched targets remain manual choices."""
	header_by_key = {header_key(header): header for header in headers if header}
	return [
		{"target_field": target, "source_column": header_by_key.get(header_key(target), "")}
		for target in IMPORT_FIELDS
	]


def _validated_mapping(import_doc, headers: list[str]) -> dict[str, str]:
	normalized_headers = [header_key(header) for header in headers if header]
	if len(normalized_headers) != len(set(normalized_headers)):
		frappe.throw(_("Başlık satırında yinelenen kolon adları var."))
	header_by_key = {header_key(header): header for header in headers}
	mapping: dict[str, str] = {}
	for row in import_doc.get("column_mappings") or []:
		target = preserve_excel_text(row.target_field)
		if is_blank(row.source_column):
			continue
		source = header_by_key.get(header_key(row.source_column))
		if target in mapping:
			frappe.throw(_("Aynı hedef alan birden fazla kez eşlenemez: {0}").format(target))
		if target not in IMPORT_FIELDS or not source:
			frappe.throw(_("Geçersiz kolon eşlemesi: {0} → {1}").format(row.source_column, target))
		mapping[target] = source
	missing = [field for field in IMPORT_FIELDS if field not in mapping]
	if missing:
		frappe.throw(_("Eksik zorunlu eşlemeler: {0}").format(", ".join(missing)))
	if len({header_key(value) for value in mapping.values()}) != len(mapping):
		frappe.throw(_("Bir Excel kolonu birden fazla hedef alana eşlenemez."))
	return mapping


def _read_rows(import_doc, content: bytes | None = None) -> list[tuple[int, dict[str, Any]]]:
	raw_rows = _read_single_worksheet(content or _get_file_content(import_doc))
	header_idx = cint(import_doc.header_row or 1) - 1
	if header_idx < 0 or header_idx >= len(raw_rows):
		frappe.throw(_("Başlık satırı Excel dosyasının dışında."))
	headers = [preserve_excel_text(value) for value in raw_rows[header_idx]]
	mapping = _validated_mapping(import_doc, headers)
	index_by_header = {header_key(value): idx for idx, value in enumerate(headers) if value}
	result = []
	for row_number, raw in enumerate(raw_rows[header_idx + 1 :], start=header_idx + 2):
		if not raw or all(is_blank(cell) for cell in raw):
			continue
		result.append(
			(
				row_number,
				{
					target: raw[index_by_header[header_key(source)]]
					if index_by_header[header_key(source)] < len(raw)
					else None
					for target, source in mapping.items()
				},
			)
		)
	return result


def build_validation_signature(import_doc, content: bytes | None = None) -> str:
	payload = {
		"content_hash": hashlib.sha256(content or _get_file_content(import_doc)).hexdigest(),
		"target_tour": import_doc.target_tour or "",
		"header_row": cint(import_doc.header_row or 1),
		"mapping": _mapping_rows(import_doc),
	}
	return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def verify_validation_signature(import_doc, content: bytes | None = None) -> None:
	if not import_doc.validation_signature or import_doc.validation_signature != build_validation_signature(
		import_doc, content
	):
		frappe.throw(
			_(
				"Dosya, hedef tur, başlık satırı veya kolon eşlemesi doğrulamadan sonra değişti. "
				"Kuru çalıştırmayı yenileyin."
			)
		)


def _find_existing_umreci(tc: str) -> str | None:
	names = frappe.get_all(
		DOCTYPE_UMRECI,
		filters={"tc_kimlik": tc},
		pluck="name",
		order_by="creation asc",
		limit_page_length=2,
	)
	if not names:
		# Compatibility for legacy identifiers saved with embedded whitespace.
		names = [
			row[0]
			for row in frappe.db.sql(
				"""
				select name
				from `tabUmreci`
				where replace(replace(replace(tc_kimlik, ' ', ''), char(9), ''), char(10), '') = %s
				order by creation asc
				limit 2
				""",
				(tc,),
			)
		]
	if len(names) > 1:
		frappe.throw(_("Aynı TC / Yabancı Kimlik ile birden fazla Umreci kaydı bulundu: {0}").format(tc))
	return names[0] if names else None


def _find_existing_booking(umreci: str, tur: str) -> str | None:
	names = frappe.get_all(
		DOCTYPE_BOOKING,
		filters={"umreci": umreci, "tur": tur},
		pluck="name",
		order_by="creation asc",
		limit_page_length=2,
	)
	if len(names) > 1:
		frappe.throw(
			_("Aynı yolcu ve tur için birden fazla rezervasyon bulundu; manuel inceleme gerekli.")
		)
	return names[0] if names else None


def _lock_import_materialization() -> None:
	"""Serialize actual import materialization to protect TC and (passenger, tour) identity."""
	locked = frappe.db.sql(
		"""
		select name
		from `tabUmre Excel Import`
		order by creation asc, name asc
		limit 1
		for update
		"""
	)
	if not locked:
		frappe.throw(_("İçe aktarım kilidi alınamadı."))


def _normalize_row(row: dict[str, Any], _tour: str) -> dict[str, Any]:
	tc = normalize_tc(row["TC KİMLİK"])
	if not tc:
		frappe.throw(_("TC / Yabancı Kimlik zorunludur."))
	arrival, return_city = split_cities(row["GELDİĞİ İL"])
	room = map_oda_tipi(row["ODA SAYISI"])
	status = normalize_status(row["YOLCU STATÜSÜ"])
	referral_text = preserve_excel_text(row.get("KİMDEN"))
	if status == "UMRECI" and not referral_text:
		frappe.throw(_("KİMDEN yalnız UMRECI statüsündeki yolcular için zorunludur."))
	price = safe_float(row["FİYAT"])
	paid = safe_float(row["ÖDEDİĞİ MİKTAR"])
	if price < 0:
		frappe.throw(_("FİYAT negatif olamaz."))
	if status == "UMRECI" and price <= 0:
		frappe.throw(_("UMRECI statüsündeki yolcular için FİYAT pozitif olmalıdır."))
	if status != "UMRECI" and price != 0:
		frappe.throw(_("UMRECI dışındaki yolcular için FİYAT sıfır olmalıdır."))
	if paid < 0:
		frappe.throw(_("ÖDEDİĞİ MİKTAR negatif olamaz."))
	if status != "UMRECI" and paid:
		frappe.throw(_("Yalnız UMRECI statüsündeki yolcular için ödenen miktar girilebilir."))
	birth = parse_date(row["DOĞUM TARİHİ"])
	if not birth:
		frappe.throw(_("DOĞUM TARİHİ geçerli bir tarih olmalıdır."))
	ad = preserve_excel_text(row["AD"])
	soyad = preserve_excel_text(row["SOYAD"])
	if not ad or not soyad:
		frappe.throw(_("AD ve SOYAD zorunludur."))
	return {
		"tc_kimlik": tc,
		"ad": ad,
		"soyad": soyad,
		"cinsiyet": normalize_cinsiyet(row["CİNSİYET"]),
		"uyruk": normalize_nationality(row["UYRUK"]),
		"dogum_tarihi": birth,
		"telefon_numarasi": sanitize_phone(row["TELEFON NUMARASI"]),
		"oda_tipi": room,
		"arrival_city": arrival,
		"return_city": return_city,
		"referral_text": referral_text,
		"statu": status,
		"ucret": price,
		"bildirilen_odenen": paid,
		"odenen": 0,
	}


def _row_key(tour: str, normalized: dict[str, Any]) -> str:
	payload = {key: value for key, value in normalized.items() if key != "referral_text"}
	payload["tour"] = tour
	return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def _booking_matches(booking, normalized: dict[str, Any], referral: str | None) -> bool:
	fields = ("oda_tipi", "arrival_city", "return_city", "statu", "ucret", "bildirilen_odenen")
	for field in fields:
		left, right = booking.get(field), normalized.get(field)
		if field in {"ucret", "bildirilen_odenen"}:
			if abs(flt(left) - flt(right)) > 0.000001:
				return False
		elif (left or "") != (right or ""):
			return False
	return (booking.get("kimden_geldi") or "") == (referral or "")


def _umreci_matches(umreci, normalized: dict[str, Any]) -> bool:
	return (
		normalize_tc(umreci.get("tc_kimlik")) == normalized["tc_kimlik"]
		and header_key(umreci.get("ad")) == header_key(normalized["ad"])
		and header_key(umreci.get("soyad")) == header_key(normalized["soyad"])
		and all(
			str(umreci.get(field) or "") == str(normalized.get(field) or "")
			for field in ("cinsiyet", "dogum_tarihi", "uyruk", "telefon_numarasi")
		)
	)


def _stage_row(
	import_doc,
	row_number: int,
	raw: dict[str, Any],
	normalized: dict[str, Any] | None,
	status: str,
	message: str | None = None,
	referral: str | None = None,
) -> dict[str, Any]:
	key = (
		_row_key(import_doc.target_tour, normalized)
		if normalized
		else hashlib.sha256(f"{import_doc.name}:{row_number}".encode()).hexdigest()
	)
	return {
		"row_number": row_number,
		"row_key": key,
		"tc_kimlik": normalized.get("tc_kimlik") if normalized else normalize_tc(raw.get("TC KİMLİK")),
		"ad": normalized.get("ad") if normalized else preserve_excel_text(raw.get("AD")),
		"soyad": normalized.get("soyad") if normalized else preserve_excel_text(raw.get("SOYAD")),
		"raw_data": json.dumps(raw, ensure_ascii=False, default=str),
		"normalized_data": json.dumps(normalized, ensure_ascii=False, default=str) if normalized else None,
		"row_status": status,
		"referral_text": normalized.get("referral_text")
		if normalized
		else preserve_excel_text(raw.get("KİMDEN")),
		"referral_source": referral,
		"umreci_link": None,
		"booking_link": None,
		"umreci_action": None,
		"booking_action": None,
		"message": message,
	}


def _resolved_referral(import_doc, normalized: dict[str, Any]) -> str | None:
	text = normalized["referral_text"]
	if not text:
		return None
	if text and frappe.db.exists(DOCTYPE_REFERRAL, text):
		return text
	if text:
		text_key = header_key(text)
		for existing in frappe.get_all(DOCTYPE_REFERRAL, pluck="name", limit_page_length=0):
			if header_key(existing) == text_key:
				return existing
	for staged in import_doc.get("staged_rows") or []:
		if (
			header_key(staged.referral_text) == header_key(text)
			and staged.referral_source
			and frappe.db.exists(DOCTYPE_REFERRAL, staged.referral_source)
		):
			return staged.referral_source
	return None


def _booking_update_is_accounting_blocked(booking) -> bool:
	"""Never mutate a snapshot already consumed by accounting or an actual payment."""
	if cost_engine._has_submitted_cost_posting(booking.name):
		return True
	if flt(booking.get("odenen") or 0) > 0:
		return True
	return bool(frappe.db.exists(
		"Umre Booking Payment",
		{
			"parent": booking.name,
			"parenttype": DOCTYPE_BOOKING,
			"parentfield": "payments",
		},
	))


def _company_from_settings() -> str:
	company = frappe.db.get_single_value("Umre Ops Settings", "company")
	if not company:
		frappe.throw(_("Umre Ops Settings üzerinde şirket seçilmelidir."))
	return company


def _apply_umreci_fields(umreci, normalized: dict[str, Any]) -> None:
	for field in ("ad", "soyad", "cinsiyet", "tc_kimlik", "dogum_tarihi", "uyruk", "telefon_numarasi"):
		umreci.set(field, normalized[field])


def _apply_booking_fields(booking, normalized: dict[str, Any], referral: str | None) -> None:
	booking.update({
		"oda_tipi": normalized["oda_tipi"],
		"kimden_geldi": referral,
		"arrival_city": normalized["arrival_city"],
		"return_city": normalized["return_city"],
		"ic_hat_baglanti": normalized["arrival_city"],
		"statu": normalized["statu"],
		"ucret": normalized["ucret"],
		"bildirilen_odenen": normalized["bildirilen_odenen"],
		"cost_policy": "System Rules",
		"cost_policy_version": "1",
		"import_row_key": _row_key(booking.tur, normalized),
		"is_imported": 1,
		"locked_financials": 1,
	})


def _validate_component_inputs_for_import(
	tour: str, normalized: dict[str, Any], existing_booking=None
) -> None:
	"""Exercise cost-rule calculations during preflight without writing documents."""
	values = {
		"tur": tour,
		"oda_tipi": normalized["oda_tipi"],
		"statu": normalized["statu"],
		"cost_policy": "System Rules",
		"yolcu_tipi": existing_booking.get("yolcu_tipi") if existing_booking else None,
		"vize_tipi": existing_booking.get("vize_tipi") if existing_booking else None,
	}
	cost_engine.validate_component_inputs(values)


def _lock_existing_booking(booking_name: str) -> None:
	if not frappe.db.sql(
		"SELECT name FROM `tabUmre Booking` WHERE name = %s FOR UPDATE", (booking_name,)
	):
		frappe.throw(_("Rezervasyon güncelleme sırasında bulunamadı: {0}").format(booking_name))


def _company_block_message(booking, settings_company: str) -> str | None:
	booking_company = booking.get("company")
	if booking_company != settings_company:
		return _("Rezervasyon şirketi ({0}) Umre Ops Settings şirketiyle ({1}) eşleşmiyor.").format(
			booking_company or _("boş"), settings_company
		)
	return None


def _merge_failed_row(rows: list[dict], failed_row: dict[str, Any]) -> list[dict]:
	merged = [row for row in rows if row.get("row_number") != failed_row.get("row_number")]
	merged.append(failed_row)
	return sorted(merged, key=lambda row: cint(row.get("row_number")))


def _process_row(
	import_doc, row_number: int, raw: dict[str, Any], summary: ImportSummary, dry_run: bool,
	company: str | None = None,
) -> dict[str, Any]:
	company = company or _company_from_settings()
	normalized = _normalize_row(raw, import_doc.target_tour)
	referral = _resolved_referral(import_doc, normalized)
	if normalized["statu"] == "UMRECI" and not referral:
		staged = _stage_row(
			import_doc,
			row_number,
			raw,
			normalized,
			"Pending Referral",
			_("Referans kaynağı eşlenmeli veya açık eylemle oluşturulmalı."),
		)
		if dry_run:
			staged["umreci_action"] = "Oluşturulacak"
			staged["booking_action"] = "Bloklu"
			return staged
		frappe.throw(staged["message"])
	existing_umreci = _find_existing_umreci(normalized["tc_kimlik"])
	if existing_umreci:
		umreci = frappe.get_doc(DOCTYPE_UMRECI, existing_umreci)
		existing_booking = _find_existing_booking(existing_umreci, import_doc.target_tour)
		if existing_booking:
			booking = frappe.get_doc(DOCTYPE_BOOKING, existing_booking)
			company_message = _company_block_message(booking, company)
			if company_message and booking.get("company"):
				staged = _stage_row(
					import_doc, row_number, raw, normalized, "Blocked Company",
					company_message, referral,
				)
			elif _booking_matches(booking, normalized, referral) and _umreci_matches(umreci, normalized):
				staged = _stage_row(import_doc, row_number, raw, normalized, "No-op", referral=referral)
			else:
				if company_message:
					staged = _stage_row(
						import_doc, row_number, raw, normalized, "Blocked Company",
						company_message, referral,
					)
				elif _booking_update_is_accounting_blocked(booking):
					staged = _stage_row(
						import_doc, row_number, raw, normalized, "Blocked Accounting",
						_(
							"Muhasebeleştirilmiş maliyet veya gerçekleşmiş ödeme nedeniyle "
							"otomatik güncelleme engellendi."
						),
						referral,
					)
				else:
					staged = _stage_row(
						import_doc, row_number, raw, normalized, "Update Ready", referral=referral
					)
			staged["umreci_link"] = existing_umreci
			staged["booking_link"] = existing_booking
			staged["umreci_action"] = (
				"Değişiklik yok" if _umreci_matches(umreci, normalized) else "Güncellenecek"
			)
			staged["booking_action"] = {
				"No-op": "Değişiklik yok", "Update Ready": "Güncellenecek",
				"Blocked Accounting": "Bloklu", "Blocked Company": "Bloklu",
			}[staged["row_status"]]
			if dry_run or staged["row_status"] == "No-op":
				if dry_run and staged["row_status"] == "Update Ready":
					_validate_component_inputs_for_import(
						import_doc.target_tour, normalized, booking
					)
				return staged
			if staged["row_status"] in {"Blocked Accounting", "Blocked Company"}:
				frappe.throw(staged["message"])
	if dry_run:
		_validate_component_inputs_for_import(import_doc.target_tour, normalized)
		staged = _stage_row(import_doc, row_number, raw, normalized, "Create Ready", referral=referral)
		staged["umreci_link"] = existing_umreci
		staged["umreci_action"] = "Değişiklik yok" if existing_umreci else "Oluşturulacak"
		staged["booking_action"] = "Oluşturulacak"
		return staged

	frappe.db.savepoint(f"excel_import_row_{row_number}")
	try:
		umreci_changed = False
		existing_booking = (
			_find_existing_booking(existing_umreci, import_doc.target_tour) if existing_umreci else None
		)
		if existing_booking:
			# Match accounting_service lock order: booking row first, then reload and guard.
			_lock_existing_booking(existing_booking)
			booking = frappe.get_doc(DOCTYPE_BOOKING, existing_booking)
			company_message = _company_block_message(booking, company)
			if company_message:
				frappe.throw(company_message)
			if _booking_update_is_accounting_blocked(booking):
				frappe.throw(
					_("Muhasebeleştirilmiş veya ödeme alınmış rezervasyon güncellenemez: {0}").format(
						booking.name
					)
				)
		if existing_umreci:
			umreci = frappe.get_doc(DOCTYPE_UMRECI, existing_umreci)
			if not _umreci_matches(umreci, normalized):
				_apply_umreci_fields(umreci, normalized)
				umreci.save()
				summary.updated_umreci += 1
				umreci_changed = True
		else:
			umreci = frappe.new_doc(DOCTYPE_UMRECI)
			_apply_umreci_fields(umreci, normalized)
			umreci.save()
			summary.created_umreci += 1
		if existing_booking:
			_apply_booking_fields(booking, normalized, referral)
			booking.flags.ignore_financial_lock = True
			booking.save()
			cost_engine.recompute_components(booking, skip_dashboard_publish=True)
			summary.updated_bookings += 1
			staged = _stage_row(import_doc, row_number, raw, normalized, "Imported", referral=referral)
			staged["umreci_link"] = umreci.name
			staged["booking_link"] = booking.name
			staged["umreci_action"] = "Güncellendi" if umreci_changed else "Değişiklik yok"
			staged["booking_action"] = "Güncellendi"
			return staged
		booking = frappe.get_doc(
			{
				"doctype": DOCTYPE_BOOKING,
				"company": company,
				"umreci": umreci.name,
				"tur": import_doc.target_tour,
				"oda_tipi": normalized["oda_tipi"],
				"kimden_geldi": referral,
				"arrival_city": normalized["arrival_city"],
				"return_city": normalized["return_city"],
				"ic_hat_baglanti": normalized["arrival_city"],
				"statu": normalized["statu"],
				"ucret": normalized["ucret"],
				"bildirilen_odenen": normalized["bildirilen_odenen"],
				"odenen": 0,
				"cost_policy": "System Rules",
				"cost_policy_version": "1",
				"import_row_key": _row_key(import_doc.target_tour, normalized),
				"is_imported": 1,
				"locked_financials": 1,
			}
		)
		booking.insert()
		summary.created_bookings += 1
		staged = _stage_row(import_doc, row_number, raw, normalized, "Imported", referral=referral)
		staged["umreci_link"] = umreci.name
		staged["booking_link"] = booking.name
		staged["umreci_action"] = (
			"Oluşturuldu" if not existing_umreci
			else ("Güncellendi" if umreci_changed else "Değişiklik yok")
		)
		staged["booking_action"] = "Oluşturuldu"
		return staged
	except Exception:
		frappe.db.rollback(save_point=f"excel_import_row_{row_number}")
		raise


def _save_result(
	import_doc,
	snapshot,
	summary: ImportSummary,
	rows: list[dict],
	status: str,
	signature: str | None = None,
	error_log: str | None = None,
) -> None:
	if snapshot:
		locked = frappe.db.sql(
			"select name from `tabUmre Excel Import` where name=%s for update", import_doc.name
		)
		if not locked:
			frappe.throw(_("İçe aktarım kaydı bulunamadı."))
		validate_dry_run_snapshot(snapshot, frappe.get_doc(DOCTYPE_IMPORT, import_doc.name))
	doc = frappe.get_doc(DOCTYPE_IMPORT, import_doc.name)
	doc.flags.ignore_import_source_guard = True
	doc.status = status
	for field, value in asdict(summary).items():
		doc.set(field, value)
	doc.set("staged_rows", [])
	for row in rows:
		doc.append("staged_rows", row)
	doc.row_log = json.dumps(rows, ensure_ascii=False, indent=2)
	doc.dry_run_result = json.dumps({"summary": asdict(summary), "rows": rows}, ensure_ascii=False, indent=2)
	doc.error_log = error_log
	if signature is not None:
		doc.validation_signature = signature
	if status in {"Validated", "Partially Completed", "Completed", "Failed"}:
		doc.completed_at = now()
	doc.save(ignore_permissions=True)
	frappe.db.commit()


def _final_status(summary: ImportSummary, *, dry_run: bool, rows: list[dict]) -> str:
	if summary.row_errors:
		return "Failed"
	if dry_run:
		return "Validated"
	return "Completed"


def run_import(docname: str, *, dry_run: bool) -> dict:
	import_doc = frappe.get_doc(DOCTYPE_IMPORT, docname)
	import_doc.check_permission("write")
	snapshot = make_source_snapshot(import_doc) if dry_run else None
	if dry_run and import_doc.status in {"Queued", "Processing", "Partially Completed", "Completed"}:
		frappe.throw(_("Bu durumda kuru çalıştırma yapılamaz."))
	if not frappe.db.exists("Umre Tour", import_doc.target_tour):
		frappe.throw(_("Hedef tur bulunamadı."))
	content = _get_file_content(import_doc)
	rows = _read_rows(import_doc, content)
	if not rows:
		frappe.throw(_("Excel dosyasında içe aktarılacak veri satırı bulunamadı."))
	try:
		company = _company_from_settings()
	except Exception as exc:
		if not dry_run:
			raise
		summary = ImportSummary(total_rows=len(rows), row_errors=len(rows) or 1)
		logs = [
			_stage_row(import_doc, row_number, raw, None, "Error", str(exc))
			for row_number, raw in rows
		]
		_save_result(
			import_doc, snapshot, summary, logs, "Failed",
			build_validation_signature(import_doc, content),
		)
		return {"summary": asdict(summary), "rows": logs}
	if not dry_run:
		verify_validation_signature(import_doc, content)
		_lock_import_materialization()
	summary = ImportSummary(total_rows=len(rows))
	logs = []
	seen_identities: set[tuple[str | None, str]] = set()
	for row_number, raw in rows:
		try:
			identity = (normalize_tc(raw.get("TC KİMLİK")), import_doc.target_tour)
			if identity in seen_identities:
				frappe.throw(_("Excel içinde aynı TC ve hedef tur birden fazla satırda yer alıyor."))
			seen_identities.add(identity)
			logs.append(_process_row(import_doc, row_number, raw, summary, dry_run, company))
		except RETRYABLE_IMPORT_ERRORS:
			raise
		except Exception as exc:
			summary.row_errors += 1
			failed_row = _stage_row(import_doc, row_number, raw, None, "Error", str(exc))
			failed_row["umreci_action"] = "Bloklu"
			failed_row["booking_action"] = "Bloklu"
			logs.append(failed_row)
			if not dry_run:
				raise ImportRowError(str(exc), failed_row) from exc
	if dry_run and any(
		row["row_status"] not in READY_ROW_STATUSES for row in logs
	):
		summary.row_errors = sum(row["row_status"] == "Error" for row in logs)
	status = _final_status(summary, dry_run=dry_run, rows=logs)
	_save_result(
		import_doc,
		snapshot,
		summary,
		logs,
		status,
		build_validation_signature(import_doc, content) if dry_run else None,
	)
	return {"summary": asdict(summary), "rows": logs}


def run_dry_run(docname: str) -> dict:
	return run_import(docname, dry_run=True)


def _mark_failed_attempt(
	docname: str, attempt_id: str, error_log: str, failed_row: dict[str, Any] | None = None
) -> None:
	frappe.db.sql("select name from `tabUmre Excel Import` where name=%s for update", docname)
	current = frappe.get_doc(DOCTYPE_IMPORT, docname)
	if not attempt_can_run(current, attempt_id):
		frappe.db.rollback()
		return
	summary = ImportSummary(total_rows=current.total_rows or 0, row_errors=(current.row_errors or 0) + 1)
	rows = json.loads(current.row_log or "[]")
	if failed_row:
		rows = _merge_failed_row(rows, failed_row)
	_save_result(current, None, summary, rows, "Failed", error_log=error_log)


def run_import_job(docname: str, attempt_id: str, user: str | None = None) -> None:
	if user:
		frappe.set_user(user)
	try:
		frappe.db.sql("select name from `tabUmre Excel Import` where name=%s for update", docname)
		doc = frappe.get_doc(DOCTYPE_IMPORT, docname)
		if not attempt_can_run(doc, attempt_id):
			frappe.db.rollback()
			return
		verify_validation_signature(doc)
		if doc.status == "Queued":
			doc.status = "Processing"
			doc.started_at = now()
			doc.completed_at = None
			doc.error_log = None
			doc.save(ignore_permissions=True)
		frappe.db.commit()
		run_import(docname, dry_run=False)
	except RETRYABLE_IMPORT_ERRORS:
		frappe.db.rollback()
		raise
	except Exception as exc:
		error_log = frappe.get_traceback()
		frappe.db.rollback()
		_mark_failed_attempt(docname, attempt_id, error_log, getattr(exc, "staged_row", None))
		frappe.db.commit()
		raise


def _enqueue_attempt(docname: str, attempt_id: str, user: str, *, after_commit: bool) -> None:
	frappe.enqueue(
		"umre_ops.umre_ops.services.excel_import_service.run_import_job",
		queue="long",
		docname=docname,
		attempt_id=attempt_id,
		user=user,
		job_id=attempt_id,
		deduplicate=True,
		enqueue_after_commit=after_commit,
	)


def enqueue_import(docname: str, user: str | None = None) -> dict:
	frappe.db.sql("select name from `tabUmre Excel Import` where name=%s for update", docname)
	doc = frappe.get_doc(DOCTYPE_IMPORT, docname)
	doc.check_permission("write")
	queue_user = user or frappe.session.user
	if doc.status in {"Queued", "Processing"} and doc.job_id:
		if not is_job_enqueued(doc.job_id):
			_enqueue_attempt(docname, doc.job_id, queue_user, after_commit=False)
		return {"job_id": doc.job_id, "status": doc.status}
	if (
		doc.status != "Validated"
		or doc.row_errors
		or cint(doc.total_rows) <= 0
		or not (doc.get("staged_rows") or [])
		or any(row.row_status not in READY_ROW_STATUSES for row in doc.get("staged_rows") or [])
	):
		frappe.throw(_("Aktarımı başlatmadan önce kuru çalıştırmayı tamamlayın."))
	verify_validation_signature(doc)
	attempt_id = uuid4().hex
	doc.status = "Queued"
	doc.job_id = attempt_id
	doc.started_at = None
	doc.completed_at = None
	doc.save(ignore_permissions=True)
	_enqueue_attempt(docname, attempt_id, queue_user, after_commit=True)
	return {"job_id": attempt_id, "status": "Queued"}


def resolve_referral(docname: str, referral_text: str, referral_source: str) -> None:
	doc = frappe.get_doc(DOCTYPE_IMPORT, docname)
	doc.check_permission("write")
	if doc.status in {"Queued", "Processing", "Partially Completed", "Completed"}:
		frappe.throw(_("Kilitli aktarım değiştirilemez."))
	if not frappe.db.exists(DOCTYPE_REFERRAL, referral_source):
		frappe.throw(_("Referans kaynağı bulunamadı."))
	for row in doc.get("staged_rows") or []:
		if header_key(row.referral_text) == header_key(referral_text):
			row.referral_source = referral_source
	doc.status = "Draft"
	doc.validation_signature = None
	doc.save()


def create_referral(docname: str, referral_text: str) -> str:
	doc = frappe.get_doc(DOCTYPE_IMPORT, docname)
	doc.check_permission("write")
	if doc.status in {"Queued", "Processing", "Partially Completed", "Completed"}:
		frappe.throw(_("Kilitli aktarım değiştirilemez."))
	if not frappe.has_permission(DOCTYPE_REFERRAL, "create", throw=False):
		frappe.throw(_("Referans kaynağı oluşturma yetkiniz yok."), frappe.PermissionError)
	name = preserve_excel_text(referral_text)
	if not name:
		frappe.throw(_("Referans adı boş olamaz."))
	existing = next(
		(
			value
			for value in frappe.get_all(DOCTYPE_REFERRAL, pluck="name", limit_page_length=0)
			if header_key(value) == header_key(name)
		),
		None,
	)
	if not existing:
		frappe.get_doc({"doctype": DOCTYPE_REFERRAL, "kaynak_adi": name}).insert()
		existing = name
	resolve_referral(docname, referral_text, existing)
	return existing
